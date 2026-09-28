"""Document parsers. Each returns a list of sections:

    {"text": str, "page": int | None, "section_title": str | None,
     "heading_hierarchy": list[str] | None}

plus page_count and warnings. Tables are converted to Markdown so row/column
relationships survive chunking. This module is imported inside the sandboxed
parser subprocess, so it must stay free of app/network dependencies.
"""

import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from backend.config import MAX_PDF_PAGES

SUPPORTED_TYPES = {
    ".pdf": ("pdf", "application/pdf"),
    ".docx": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ".txt": ("text", "text/plain"),
    ".md": ("markdown", "text/markdown"),
    ".markdown": ("markdown", "text/markdown"),
    ".epub": ("epub", "application/epub+zip"),
}


class UnsupportedDocument(ValueError):
    pass


def detect_type(filename: str, content: bytes) -> tuple[str, str, str]:
    """Validate extension against magic bytes. Returns (kind, extension, mime_type)."""
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_TYPES:
        shown = ext or "none"
        raise UnsupportedDocument(f"Unsupported file type {shown!r}. Supported: {', '.join(SUPPORTED_TYPES)}")
    kind, mime = SUPPORTED_TYPES[ext]

    if kind == "pdf" and not content.startswith(b"%PDF-"):
        raise UnsupportedDocument("File content is not a valid PDF")
    if kind in ("docx", "epub") and not content.startswith(b"PK\x03\x04"):
        raise UnsupportedDocument(f"File content is not a valid {ext} archive")
    if kind in ("text", "markdown") and b"\x00" in content[:8192]:
        raise UnsupportedDocument("Text file appears to be binary")
    return kind, ext, mime


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _clean_cell(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().replace("|", "\\|")


def table_to_markdown(rows: List[List[Any]]) -> str:
    rows = [[_clean_cell(c) for c in row] for row in rows if row and any(c for c in row)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, body = rows[0], rows[1:]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in body]
    return "\n".join(lines)


def _section(text: str, page: Optional[int] = None, title: Optional[str] = None,
             hierarchy: Optional[List[str]] = None) -> Dict[str, Any]:
    return {"text": text.strip(), "page": page, "section_title": title, "heading_hierarchy": hierarchy or None}


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def parse_pdf(path: str) -> Dict[str, Any]:
    import pdfplumber

    sections, warnings = [], []
    with pdfplumber.open(path) as pdf:
        page_count = len(pdf.pages)
        if page_count > MAX_PDF_PAGES:
            raise UnsupportedDocument(f"PDF has {page_count} pages; the limit is {MAX_PDF_PAGES}")

        for page_no, page in enumerate(pdf.pages, start=1):
            tables = page.find_tables()
            bboxes = [t.bbox for t in tables]

            def outside_tables(obj, _bboxes=bboxes):
                if obj.get("object_type") != "char":
                    return True
                x, top = obj["x0"], obj["top"]
                return not any(b[0] <= x <= b[2] and b[1] <= top <= b[3] for b in _bboxes)

            text = (page.filter(outside_tables).extract_text() or "") if bboxes else (page.extract_text() or "")
            table_md = [md for md in (table_to_markdown(t.extract()) for t in tables) if md]
            body = "\n\n".join([text.strip()] + table_md).strip()

            if body:
                sections.append(_section(body, page=page_no))
            elif page.images:
                warnings.append(
                    f"Page {page_no} has no extractable text (scanned image?); OCR is not enabled, so it was skipped"
                )
            getattr(page, "flush_cache", lambda: None)()  # release per-page memory

    return {"sections": sections, "page_count": page_count, "warnings": warnings}


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def parse_docx(path: str) -> Dict[str, Any]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(path)
    sections: List[Dict[str, Any]] = []
    hierarchy: List[Optional[str]] = [None] * 6
    parts: List[str] = []
    title: Optional[str] = None

    def flush():
        text = "\n\n".join(p for p in parts if p).strip()
        if text:
            sections.append(_section(text, title=title, hierarchy=[h for h in hierarchy if h]))

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            style = (para.style.name if para.style is not None else "") or ""
            match = re.match(r"Heading (\d)", style)
            if match and text:
                flush()
                parts = []
                level = min(int(match.group(1)), 6)
                hierarchy[level - 1] = text
                for i in range(level, 6):
                    hierarchy[i] = None
                title = text
            elif text:
                parts.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            md = table_to_markdown([[cell.text for cell in row.cells] for row in table.rows])
            if md:
                parts.append(md)
    flush()
    return {"sections": sections, "page_count": None, "warnings": []}


# ---------------------------------------------------------------------------
# Plain text / Markdown
# ---------------------------------------------------------------------------

def _decode(data: bytes, warnings: List[str]) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        warnings.append("File is not valid UTF-8; decoded as Windows-1252")
        return data.decode("cp1252", errors="replace")


def parse_text(path: str, markdown: bool) -> Dict[str, Any]:
    warnings: List[str] = []
    text = _decode(Path(path).read_bytes(), warnings)
    if not markdown:
        return {"sections": [_section(text)] if text.strip() else [], "page_count": None, "warnings": warnings}

    sections: List[Dict[str, Any]] = []
    hierarchy: List[Optional[str]] = [None] * 6
    title: Optional[str] = None
    buf: List[str] = []
    in_code = False

    def flush():
        body = "\n".join(buf).strip()
        if body:
            sections.append(_section(body, title=title, hierarchy=[h for h in hierarchy if h]))

    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            in_code = not in_code
        heading = None if in_code else re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            flush()
            buf = []
            level = len(heading.group(1))
            title = heading.group(2)
            hierarchy[level - 1] = title
            for i in range(level, 6):
                hierarchy[i] = None
        else:
            buf.append(line)
    flush()
    return {"sections": sections, "page_count": None, "warnings": warnings}


# ---------------------------------------------------------------------------
# EPUB
# ---------------------------------------------------------------------------

def parse_epub(path: str, filename: str) -> Dict[str, Any]:
    from backend.ingestion.epub_processor import EPUBProcessor

    docs = EPUBProcessor().process(path)
    sections = [
        _section(d["text"], title=d.get("section_title") or d.get("chapter_title"),
                 hierarchy=d.get("heading_hierarchy"))
        for d in docs if d.get("text", "").strip()
    ]
    return {"sections": sections, "page_count": None, "warnings": []}


def parse_file(kind: str, path: str, filename: str) -> Dict[str, Any]:
    if kind == "pdf":
        return parse_pdf(path)
    if kind == "docx":
        return parse_docx(path)
    if kind in ("text", "markdown"):
        return parse_text(path, markdown=kind == "markdown")
    if kind == "epub":
        return parse_epub(path, filename)
    raise UnsupportedDocument(f"No parser for {kind}")
